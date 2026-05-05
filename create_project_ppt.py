"""Build animated, branded PPT for AI Ticket Agent (IntelliTriage) project."""
from pathlib import Path
from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn

REPO = Path(__file__).parent
BG_IMG = REPO / "ppt_media_dump" / "image1.jpeg"
OUT_FILE = REPO / "AI-Ticket-Agent-Project-Deck.pptx"

C_BG_DARK = RGBColor(0x10, 0x12, 0x3A)
C_PANEL = RGBColor(0x1A, 0x1F, 0x55)
C_PANEL_LITE = RGBColor(0x26, 0x2D, 0x73)
C_ACCENT = RGBColor(0x5E, 0xC8, 0xFF)
C_ACCENT_2 = RGBColor(0xA8, 0x7B, 0xFF)
C_ACCENT_3 = RGBColor(0x35, 0xE0, 0xC1)
C_ACCENT_4 = RGBColor(0xFF, 0xB1, 0x4A)
C_TEXT = RGBColor(0xF4, 0xF7, 0xFF)
C_TEXT_DIM = RGBColor(0xC8, 0xD2, 0xF2)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
TOTAL = 11


# ── Animation helpers ─────────────────────────────────────────────────────────
def _ensure_timing(slide):
    sld = slide._element
    existing = sld.find(qn("p:timing"))
    if existing is not None:
        sld.remove(existing)
    timing_xml = """
<p:timing xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">
  <p:tnLst>
    <p:par>
      <p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot">
        <p:childTnLst>
          <p:seq concurrent="1" nextAc="seek">
            <p:cTn id="2" dur="indefinite" nodeType="mainSeq">
              <p:childTnLst/>
            </p:cTn>
            <p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>
            <p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst>
          </p:seq>
        </p:childTnLst>
      </p:cTn>
    </p:par>
  </p:tnLst>
</p:timing>
"""
    timing = etree.fromstring(timing_xml)
    sld.append(timing)
    ns = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main"}
    seq_main = timing.find(".//p:cTn[@nodeType='mainSeq']", ns)
    return seq_main.find("p:childTnLst", ns)


def _fade_in_par(shape_id: int, delay_ms: int) -> str:
    return f"""
<p:par xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">
  <p:cTn id="{shape_id*10}" presetID="10" presetClass="entr" presetSubtype="0" fill="hold" nodeType="afterEffect">
    <p:stCondLst><p:cond delay="{delay_ms}"/></p:stCondLst>
    <p:childTnLst>
      <p:set>
        <p:cBhvr>
          <p:cTn id="{shape_id*10+1}" dur="1" fill="hold">
            <p:stCondLst><p:cond delay="0"/></p:stCondLst>
          </p:cTn>
          <p:tgtEl><p:spTgt spid="{shape_id}"/></p:tgtEl>
          <p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst>
        </p:cBhvr>
        <p:to><p:strVal val="visible"/></p:to>
      </p:set>
      <p:anim calcmode="lin" valueType="num">
        <p:cBhvr additive="base">
          <p:cTn id="{shape_id*10+2}" dur="500" fill="hold"/>
          <p:tgtEl><p:spTgt spid="{shape_id}"/></p:tgtEl>
          <p:attrNameLst><p:attrName>style.opacity</p:attrName></p:attrNameLst>
        </p:cBhvr>
        <p:tavLst>
          <p:tav tm="0"><p:val><p:fltVal val="0"/></p:val></p:tav>
          <p:tav tm="100000"><p:val><p:fltVal val="1"/></p:val></p:tav>
        </p:tavLst>
      </p:anim>
    </p:childTnLst>
  </p:cTn>
</p:par>
"""


def add_entrance(slide, shapes, stagger_ms: int = 220):
    seq = _ensure_timing(slide)
    for i, shp in enumerate(shapes):
        seq.append(etree.fromstring(_fade_in_par(shp.shape_id, i * stagger_ms)))


# ── Visual helpers ────────────────────────────────────────────────────────────
def add_background(slide):
    pic = slide.shapes.add_picture(str(BG_IMG), 0, 0, width=SW, height=SH)
    spTree = pic._element.getparent()
    spTree.remove(pic._element)
    spTree.insert(2, pic._element)
    overlay = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    overlay.line.fill.background()
    overlay.fill.solid()
    overlay.fill.fore_color.rgb = C_BG_DARK
    sp = overlay.fill.fore_color._xFill
    alpha = etree.SubElement(sp, qn("a:alpha"))
    alpha.set("val", "55000")
    return overlay


def add_text(slide, text, left, top, width, height, size=18, color=C_TEXT, bold=False,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    lines = text.split("\n") if isinstance(text, str) else list(text)
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        run.font.name = "Segoe UI"
    return box


def add_panel(slide, left, top, width, height, fill=C_PANEL, line=C_ACCENT, line_w=1.0):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    sp = shape.fill.fore_color._xFill
    alpha = etree.SubElement(sp, qn("a:alpha"))
    alpha.set("val", "82000")
    shape.line.color.rgb = line
    shape.line.width = Pt(line_w)
    return shape


def add_icon_chip(slide, glyph, left, top, size=Inches(0.85), bg=C_ACCENT, fg=C_BG_DARK):
    chip = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, size, size)
    chip.fill.solid()
    chip.fill.fore_color.rgb = bg
    chip.line.color.rgb = bg
    tf = chip.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = glyph
    r.font.size = Pt(22)
    r.font.bold = True
    r.font.color.rgb = fg
    r.font.name = "Segoe UI Symbol"
    return chip


def add_brand_strip(slide):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, Inches(0.08))
    bar.fill.solid()
    bar.fill.fore_color.rgb = C_ACCENT
    bar.line.fill.background()


def add_footer(slide, page_num):
    add_text(slide, "IntelliTriage • AI-Powered Internal Ticketing",
             Inches(0.4), Inches(7.05), Inches(8), Inches(0.35),
             size=10, color=C_TEXT_DIM)
    add_text(slide, f"{page_num} / {TOTAL}",
             Inches(12.2), Inches(7.05), Inches(0.9), Inches(0.35),
             size=10, color=C_TEXT_DIM, align=PP_ALIGN.RIGHT)


def slide_blank():
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(s)
    add_brand_strip(s)
    return s


def arrow(slide, left, top, w, h, color=C_ACCENT, shape=MSO_SHAPE.RIGHT_ARROW):
    a = slide.shapes.add_shape(shape, left, top, w, h)
    a.fill.solid(); a.fill.fore_color.rgb = color
    a.line.color.rgb = color
    return a


# ── Slides ────────────────────────────────────────────────────────────────────
def slide_title():
    s = slide_blank()
    chip = add_icon_chip(s, "🤖", Inches(0.9), Inches(2.5), size=Inches(1.2), bg=C_ACCENT)
    eyebrow = add_text(s, "AI-POWERED INTERNAL SUPPORT",
                       Inches(2.4), Inches(2.55), Inches(8), Inches(0.5),
                       size=14, color=C_ACCENT, bold=True)
    title = add_text(s, "IntelliTriage",
                     Inches(2.4), Inches(2.95), Inches(10), Inches(1.4),
                     size=72, color=C_TEXT, bold=True)
    sub = add_text(s, "Smart clarification • Auto routing • Live insights",
                   Inches(2.4), Inches(4.4), Inches(10), Inches(0.7),
                   size=24, color=C_TEXT_DIM)
    tagline_panel = add_panel(s, Inches(2.4), Inches(5.4), Inches(8.5), Inches(0.9),
                              fill=C_PANEL_LITE, line=C_ACCENT_2, line_w=1.2)
    add_text(s, "Project overview · Architecture · Benefits · Roadmap",
             Inches(2.6), Inches(5.55), Inches(8.2), Inches(0.7),
             size=18, color=C_TEXT, align=PP_ALIGN.CENTER)
    add_footer(s, 1)
    add_entrance(s, [chip, eyebrow, title, sub, tagline_panel], stagger_ms=300)


def slide_executive():
    s = slide_blank()
    add_text(s, "🚀  Executive Summary",
             Inches(0.6), Inches(0.4), Inches(12), Inches(0.7),
             size=34, color=C_TEXT, bold=True)
    add_text(s, "Why IntelliTriage exists, and what it delivers",
             Inches(0.6), Inches(1.1), Inches(12), Inches(0.5),
             size=16, color=C_TEXT_DIM)
    items = [
        ("🎯", "Mission",     "Reduce manual triage by routing every ticket with AI-grade precision."),
        ("⚡", "Speed",        "Sub-second AI response with priority-aware processing."),
        ("🧠", "Intelligence", "Asks clarifying questions only when context is missing."),
        ("📊", "Visibility",   "Live analytics for teams, urgency, and trends."),
    ]
    panels = []
    x0, y0 = Inches(0.6), Inches(2.0)
    w, h, gap = Inches(3.0), Inches(4.6), Inches(0.15)
    for i, (icon, title, body) in enumerate(items):
        left = x0 + i * (w + gap)
        panel = add_panel(s, left, y0, w, h)
        add_icon_chip(s, icon, left + Inches(0.25), y0 + Inches(0.3))
        add_text(s, title, left + Inches(1.25), y0 + Inches(0.4), w - Inches(1.4), Inches(0.6),
                 size=22, color=C_TEXT, bold=True)
        add_text(s, body, left + Inches(0.3), y0 + Inches(1.6), w - Inches(0.6), Inches(2.5),
                 size=15, color=C_TEXT_DIM)
        panels.append(panel)
    add_footer(s, 2)
    add_entrance(s, panels, stagger_ms=220)


def slide_architecture():
    s = slide_blank()
    add_text(s, "🏗  Solution Architecture",
             Inches(0.6), Inches(0.4), Inches(12), Inches(0.7),
             size=34, color=C_TEXT, bold=True)
    add_text(s, "End-to-end flow from user to AI to data",
             Inches(0.6), Inches(1.1), Inches(12), Inches(0.5),
             size=16, color=C_TEXT_DIM)

    user_panel = add_panel(s, Inches(0.5), Inches(2.4), Inches(2.4), Inches(2.0), line=C_ACCENT)
    add_icon_chip(s, "👤", Inches(0.7), Inches(2.55), size=Inches(0.7))
    add_text(s, "User", Inches(1.5), Inches(2.6), Inches(1.5), Inches(0.5),
             size=20, color=C_TEXT, bold=True)
    add_text(s, "Submits issue\nText or voice\nReceives clarifications",
             Inches(0.7), Inches(3.3), Inches(2.0), Inches(1.1),
             size=12, color=C_TEXT_DIM)

    fe_panel = add_panel(s, Inches(3.3), Inches(2.4), Inches(2.4), Inches(2.0), line=C_ACCENT_2)
    add_icon_chip(s, "🖥", Inches(3.5), Inches(2.55), size=Inches(0.7), bg=C_ACCENT_2)
    add_text(s, "Frontend", Inches(4.3), Inches(2.6), Inches(2.0), Inches(0.5),
             size=20, color=C_TEXT, bold=True)
    add_text(s, "React + Vite\nDashboards, voice\nNginx container",
             Inches(3.5), Inches(3.3), Inches(2.0), Inches(1.1),
             size=12, color=C_TEXT_DIM)

    be_panel = add_panel(s, Inches(6.1), Inches(2.4), Inches(2.4), Inches(2.0), line=C_ACCENT_3)
    add_icon_chip(s, "⚙", Inches(6.3), Inches(2.55), size=Inches(0.7), bg=C_ACCENT_3)
    add_text(s, "Backend API", Inches(7.1), Inches(2.6), Inches(2.5), Inches(0.5),
             size=20, color=C_TEXT, bold=True)
    add_text(s, "FastAPI + SQLAlchemy\nAuth · Chat · Tickets\nClarity guardrails",
             Inches(6.3), Inches(3.3), Inches(2.0), Inches(1.1),
             size=12, color=C_TEXT_DIM)

    ai_panel = add_panel(s, Inches(8.9), Inches(0.85), Inches(3.9), Inches(2.0), line=C_ACCENT_4)
    add_icon_chip(s, "🧠", Inches(9.1), Inches(1.0), size=Inches(0.7), bg=C_ACCENT_4)
    add_text(s, "HCL AICafe (LLM)", Inches(9.9), Inches(1.05), Inches(2.8), Inches(0.5),
             size=20, color=C_TEXT, bold=True)
    add_text(s, "Clarification agent\nClassification agent\nJSON-structured output",
             Inches(9.1), Inches(1.75), Inches(3.5), Inches(1.1),
             size=12, color=C_TEXT_DIM)

    db_panel = add_panel(s, Inches(8.9), Inches(3.0), Inches(3.9), Inches(2.0), line=C_ACCENT)
    add_icon_chip(s, "🗄", Inches(9.1), Inches(3.15), size=Inches(0.7))
    add_text(s, "Data Layer (SQLite)", Inches(9.9), Inches(3.2), Inches(2.8), Inches(0.5),
             size=20, color=C_TEXT, bold=True)
    add_text(s, "Users · Tickets · Feedback\nLocal/demo persistence",
             Inches(9.1), Inches(3.9), Inches(3.5), Inches(1.0),
             size=12, color=C_TEXT_DIM)

    a1 = arrow(s, Inches(2.95), Inches(3.25), Inches(0.35), Inches(0.3))
    a2 = arrow(s, Inches(5.75), Inches(3.25), Inches(0.35), Inches(0.3), color=C_ACCENT_2)
    up = arrow(s, Inches(7.85), Inches(2.0), Inches(0.4), Inches(0.6),
               color=C_ACCENT_4, shape=MSO_SHAPE.UP_ARROW)
    dn = arrow(s, Inches(8.4), Inches(2.0), Inches(0.4), Inches(0.6),
               color=C_ACCENT_3, shape=MSO_SHAPE.DOWN_ARROW)
    a3 = arrow(s, Inches(8.55), Inches(3.85), Inches(0.4), Inches(0.3))

    cap = add_panel(s, Inches(0.5), Inches(5.3), Inches(12.3), Inches(1.4),
                    fill=C_PANEL, line=C_ACCENT_2)
    add_text(s,
             "Flow:  User submits request  ➜  Frontend posts to /chat  ➜  Backend evaluates clarity\n"
             "          ➜  AI asks clarification OR Backend creates ticket  ➜  Persisted in SQLite & shown in dashboard",
             Inches(0.7), Inches(5.45), Inches(12.0), Inches(1.2),
             size=15, color=C_TEXT)

    add_footer(s, 3)
    add_entrance(s, [user_panel, a1, fe_panel, a2, be_panel, up, dn, ai_panel, a3, db_panel, cap],
                 stagger_ms=180)


def slide_decision_flow():
    s = slide_blank()
    add_text(s, "🧭  AI Decision Flow",
             Inches(0.6), Inches(0.4), Inches(12), Inches(0.7),
             size=34, color=C_TEXT, bold=True)
    add_text(s, "Asks only when needed. Creates the ticket the moment clarity is met.",
             Inches(0.6), Inches(1.1), Inches(12), Inches(0.5),
             size=16, color=C_TEXT_DIM)

    steps = [
        ("📨", "User submits request", C_ACCENT),
        ("🔍", "Backend extracts clarity signals\n(system, issue, impact/urgency)", C_ACCENT_2),
        ("❓", "Missing details?\nAsk one targeted clarification", C_ACCENT_4),
        ("✅", "Complete?\nCreate INC ticket + assign team & priority", C_ACCENT_3),
    ]
    panels = []
    x0, y0 = Inches(0.5), Inches(2.2)
    w, h = Inches(2.95), Inches(3.4)
    for i, (icon, body, accent) in enumerate(steps):
        left = x0 + i * (w + Inches(0.18))
        panel = add_panel(s, left, y0, w, h, line=accent, line_w=1.4)
        add_icon_chip(s, icon, left + (w - Inches(0.95)) / 2, y0 + Inches(0.4),
                      size=Inches(0.95), bg=accent)
        add_text(s, f"Step {i+1}", left, y0 + Inches(1.5), w, Inches(0.45),
                 size=14, color=accent, bold=True, align=PP_ALIGN.CENTER)
        add_text(s, body, left + Inches(0.2), y0 + Inches(2.0), w - Inches(0.4), Inches(1.3),
                 size=14, color=C_TEXT, align=PP_ALIGN.CENTER)
        panels.append(panel)
        if i < len(steps) - 1:
            arrow(s, left + w, y0 + Inches(1.55), Inches(0.18), Inches(0.3))

    add_text(s,
             "Guardrail: even if the model says “no clarification needed”, the backend re-checks for\n"
             "system + issue + impact/urgency before creating the ticket.",
             Inches(0.6), Inches(6.0), Inches(12), Inches(0.9),
             size=14, color=C_TEXT_DIM)

    add_footer(s, 4)
    add_entrance(s, panels, stagger_ms=240)


def slide_features():
    s = slide_blank()
    add_text(s, "✨  Key Features",
             Inches(0.6), Inches(0.4), Inches(12), Inches(0.7),
             size=34, color=C_TEXT, bold=True)
    add_text(s, "Capabilities packed into the daily user experience",
             Inches(0.6), Inches(1.1), Inches(12), Inches(0.5),
             size=16, color=C_TEXT_DIM)
    features = [
        ("🤖", "AI Clarification Agent", "Targeted follow-up questions"),
        ("🎯", "Smart Routing",          "Team + P1–P4 priority"),
        ("🎙", "Voice Assistant",         "Hands-free ticket capture"),
        ("📊", "Live Dashboard",          "Volume, trends, priorities"),
        ("🔐", "Secure Auth",             "JWT + email verification"),
        ("💬", "AI Feedback Loop",        "Rate AI for continual improvement"),
    ]
    panels = []
    cols = 3
    w, h = Inches(4.0), Inches(2.45)
    gx, gy = Inches(0.2), Inches(0.25)
    x0, y0 = Inches(0.55), Inches(2.0)
    for i, (icon, title, body) in enumerate(features):
        c = i % cols; r = i // cols
        left = x0 + c * (w + gx); top = y0 + r * (h + gy)
        panel = add_panel(s, left, top, w, h)
        add_icon_chip(s, icon, left + Inches(0.25), top + Inches(0.3),
                      size=Inches(0.8), bg=C_ACCENT_2)
        add_text(s, title, left + Inches(1.25), top + Inches(0.35), w - Inches(1.4), Inches(0.6),
                 size=20, color=C_TEXT, bold=True)
        add_text(s, body, left + Inches(0.3), top + Inches(1.3), w - Inches(0.6), Inches(1.0),
                 size=14, color=C_TEXT_DIM)
        panels.append(panel)
    add_footer(s, 5)
    add_entrance(s, panels, stagger_ms=160)


def slide_benefits():
    s = slide_blank()
    add_text(s, "🌟  Business Benefits",
             Inches(0.6), Inches(0.4), Inches(12), Inches(0.7),
             size=34, color=C_TEXT, bold=True)
    add_text(s, "Operational gains and improved user experience",
             Inches(0.6), Inches(1.1), Inches(12), Inches(0.5),
             size=16, color=C_TEXT_DIM)

    panels = []
    left_panel = add_panel(s, Inches(0.55), Inches(2.0), Inches(6.0), Inches(4.6), line=C_ACCENT)
    add_icon_chip(s, "🏢", Inches(0.75), Inches(2.18), size=Inches(0.75))
    add_text(s, "For Operations", Inches(1.65), Inches(2.25), Inches(5), Inches(0.6),
             size=24, color=C_TEXT, bold=True)
    add_text(s,
             "✓  Cleaner tickets, fewer back-and-forths\n"
             "✓  Consistent team assignment\n"
             "✓  Automatic priority tagging (P1–P4)\n"
             "✓  Reduced manual triage time\n"
             "✓  Audit trail across user + AI exchanges",
             Inches(0.85), Inches(3.1), Inches(5.5), Inches(3.5),
             size=16, color=C_TEXT)
    panels.append(left_panel)

    right_panel = add_panel(s, Inches(6.85), Inches(2.0), Inches(6.0), Inches(4.6), line=C_ACCENT_3)
    add_icon_chip(s, "🙋", Inches(7.05), Inches(2.18), size=Inches(0.75), bg=C_ACCENT_3)
    add_text(s, "For Users", Inches(7.95), Inches(2.25), Inches(5), Inches(0.6),
             size=24, color=C_TEXT, bold=True)
    add_text(s,
             "✓  Guided ticket creation\n"
             "✓  Faster first response\n"
             "✓  Voice + text accessibility\n"
             "✓  Real-time status visibility\n"
             "✓  Clear progress through the AI flow",
             Inches(7.15), Inches(3.1), Inches(5.5), Inches(3.5),
             size=16, color=C_TEXT)
    panels.append(right_panel)

    add_footer(s, 6)
    add_entrance(s, panels, stagger_ms=300)


def slide_apis():
    s = slide_blank()
    add_text(s, "🔌  API & Platform Surface",
             Inches(0.6), Inches(0.4), Inches(12), Inches(0.7),
             size=34, color=C_TEXT, bold=True)
    add_text(s, "Modular, REST-style, container-ready",
             Inches(0.6), Inches(1.1), Inches(12), Inches(0.5),
             size=16, color=C_TEXT_DIM)

    groups = [
        ("🩺  Health",  ["GET /health/aicafe"]),
        ("💬  Chat",    ["POST /chat", "GET /chat (help)"]),
        ("🎫  Tickets", ["GET /tickets", "PUT /tickets/{id}",
                         "POST /tickets/{id}/cancel", "POST /tickets/{id}/feedback"]),
        ("🔐  Auth",    ["POST /auth/register", "POST /auth/login",
                         "POST /auth/forgot-password", "POST /auth/reset-password",
                         "GET /auth/me"]),
        ("👮  Admin",   ["GET /admin/feedbacks"]),
        ("📦  Runtime", ["Frontend :3000 (Nginx)", "Backend :8080 (Uvicorn)",
                         "Podman bridge network"]),
    ]
    panels = []
    cols = 3
    w, h = Inches(4.0), Inches(2.4)
    gx, gy = Inches(0.2), Inches(0.25)
    x0, y0 = Inches(0.55), Inches(2.0)
    for i, (heading, lines) in enumerate(groups):
        c = i % cols; r = i // cols
        left = x0 + c * (w + gx); top = y0 + r * (h + gy)
        panel = add_panel(s, left, top, w, h, line=C_ACCENT_2)
        add_text(s, heading, left + Inches(0.25), top + Inches(0.2), w - Inches(0.5), Inches(0.5),
                 size=18, color=C_ACCENT, bold=True)
        add_text(s, "\n".join(lines),
                 left + Inches(0.3), top + Inches(0.85), w - Inches(0.5), h - Inches(1.0),
                 size=13, color=C_TEXT)
        panels.append(panel)
    add_footer(s, 7)
    add_entrance(s, panels, stagger_ms=140)


def slide_security():
    s = slide_blank()
    add_text(s, "🛡  Security & Reliability",
             Inches(0.6), Inches(0.4), Inches(12), Inches(0.7),
             size=34, color=C_TEXT, bold=True)
    add_text(s, "Defense-in-depth with graceful degradation",
             Inches(0.6), Inches(1.1), Inches(12), Inches(0.5),
             size=16, color=C_TEXT_DIM)
    items = [
        ("🔐", "JWT Auth + RBAC",   "Tokens, roles, signed sessions"),
        ("🧂", "Password Security", "Bcrypt hashing, reset & verify flows"),
        ("🌐", "CORS Hardening",    "Strict allow-list of origins"),
        ("🩺", "Health Telemetry",  "Live AI reachability surfacing"),
        ("🛟", "Graceful Failure",  "503s with actionable messages"),
        ("📜", "Audit Trail",       "Tickets retain original + clarified text"),
    ]
    panels = []
    cols = 3
    w, h = Inches(4.0), Inches(2.4)
    gx, gy = Inches(0.2), Inches(0.25)
    x0, y0 = Inches(0.55), Inches(2.0)
    for i, (icon, title, body) in enumerate(items):
        c = i % cols; r = i // cols
        left = x0 + c * (w + gx); top = y0 + r * (h + gy)
        panel = add_panel(s, left, top, w, h, line=C_ACCENT_3)
        add_icon_chip(s, icon, left + Inches(0.25), top + Inches(0.3),
                      size=Inches(0.8), bg=C_ACCENT_3)
        add_text(s, title, left + Inches(1.25), top + Inches(0.35), w - Inches(1.4), Inches(0.6),
                 size=18, color=C_TEXT, bold=True)
        add_text(s, body, left + Inches(0.3), top + Inches(1.3), w - Inches(0.6), Inches(1.0),
                 size=13, color=C_TEXT_DIM)
        panels.append(panel)
    add_footer(s, 8)
    add_entrance(s, panels, stagger_ms=140)


def slide_tech_stack():
    s = slide_blank()
    add_text(s, "🧰  Technology Stack",
             Inches(0.6), Inches(0.4), Inches(12), Inches(0.7),
             size=34, color=C_TEXT, bold=True)
    add_text(s, "Modern, lightweight, container-friendly",
             Inches(0.6), Inches(1.1), Inches(12), Inches(0.5),
             size=16, color=C_TEXT_DIM)
    rows = [
        ("Frontend", "React 18 · Vite · Material Symbols · Speech API", C_ACCENT_2),
        ("Backend",  "FastAPI · Pydantic · SQLAlchemy · httpx",        C_ACCENT_3),
        ("AI Layer", "HCL AICafe (OpenAI-compatible) · gpt-4.1",       C_ACCENT_4),
        ("Storage",  "SQLite (local/demo)",                            C_ACCENT),
        ("Email",    "Brevo SMTP relay (verify + reset)",              C_ACCENT_2),
        ("Runtime",  "Podman containers · Bridge network",             C_ACCENT_3),
    ]
    panels = []
    y = Inches(2.0)
    for label, body, color in rows:
        panel = add_panel(s, Inches(0.55), y, Inches(12.2), Inches(0.7), line=color)
        add_text(s, label, Inches(0.85), y + Inches(0.15), Inches(2.2), Inches(0.4),
                 size=18, color=color, bold=True)
        add_text(s, body, Inches(3.2), y + Inches(0.15), Inches(9.4), Inches(0.4),
                 size=15, color=C_TEXT)
        panels.append(panel)
        y += Inches(0.78)
    add_footer(s, 9)
    add_entrance(s, panels, stagger_ms=150)


def slide_roadmap():
    s = slide_blank()
    add_text(s, "🗺  Roadmap",
             Inches(0.6), Inches(0.4), Inches(12), Inches(0.7),
             size=34, color=C_TEXT, bold=True)
    add_text(s, "Where IntelliTriage is heading next",
             Inches(0.6), Inches(1.1), Inches(12), Inches(0.5),
             size=16, color=C_TEXT_DIM)
    phases = [
        ("Now",   "Clarification + routing\nDashboard analytics\nVoice assistant", C_ACCENT),
        ("Next",  "Confidence scoring\nMulti-language support\nReport exports",     C_ACCENT_2),
        ("Later", "Pluggable AI providers\nProduction DB + observability\nWorkflow automation", C_ACCENT_3),
    ]
    panels = []
    x0, y0 = Inches(0.6), Inches(2.1)
    w, h, gap = Inches(4.0), Inches(4.4), Inches(0.18)
    for i, (label, body, color) in enumerate(phases):
        left = x0 + i * (w + gap)
        panel = add_panel(s, left, y0, w, h, line=color, line_w=1.4)
        add_text(s, label, left, y0 + Inches(0.35), w, Inches(0.7),
                 size=26, color=color, bold=True, align=PP_ALIGN.CENTER)
        add_text(s, body,
                 left + Inches(0.3), y0 + Inches(1.45), w - Inches(0.6), h - Inches(1.6),
                 size=18, color=C_TEXT, align=PP_ALIGN.CENTER)
        panels.append(panel)
        if i < len(phases) - 1:
            arrow(s, left + w, y0 + Inches(2.0), Inches(0.18), Inches(0.3), color=color)
    add_footer(s, 10)
    add_entrance(s, panels, stagger_ms=300)


def slide_thanks():
    s = slide_blank()
    add_icon_chip(s, "🤖", Inches(5.9), Inches(1.6), size=Inches(1.4))
    add_text(s, "Thank You",
             Inches(0.5), Inches(3.2), Inches(12.3), Inches(1.4),
             size=72, color=C_TEXT, bold=True, align=PP_ALIGN.CENTER)
    add_text(s, "Questions, feedback, demos — let's talk.",
             Inches(0.5), Inches(4.5), Inches(12.3), Inches(0.6),
             size=22, color=C_TEXT_DIM, align=PP_ALIGN.CENTER)
    panel = add_panel(s, Inches(3.4), Inches(5.4), Inches(6.5), Inches(1.0),
                      fill=C_PANEL, line=C_ACCENT_2)
    add_text(s, "🌐  Demo: http://localhost:3000     ⚙  API: http://localhost:8080",
             Inches(3.6), Inches(5.6), Inches(6.1), Inches(0.6),
             size=16, color=C_TEXT, align=PP_ALIGN.CENTER)
    add_footer(s, 11)
    add_entrance(s, [panel], stagger_ms=200)


slide_title()
slide_executive()
slide_architecture()
slide_decision_flow()
slide_features()
slide_benefits()
slide_apis()
slide_security()
slide_tech_stack()
slide_roadmap()
slide_thanks()

prs.save(OUT_FILE)
print(f"Created: {OUT_FILE}")
